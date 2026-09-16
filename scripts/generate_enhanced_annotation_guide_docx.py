#!/usr/bin/env python3
"""
Generate the Comprehensive, Publication-Grade Legal Annotation Guide (.docx).

Includes:
- Official Institutional Header & Thesis Metadata
- Comprehensive Introduction & AI System Architecture Context
- Data Privacy, Anonymity & Research Ethics (RA 10173)
- Pairwise Evaluation Scope & The Single-Premise Rule
- The Three Evaluation Categories (Contradiction, Entailment, Neutral)
- Critical Legal Advisory: Local Autonomy vs. Ultra Vires (Magtajas Doctrine)
- The Three Difficulty Tiers & Evaluator Calibration
- Evaluator Confidence Rating (1-3) & Legal Notes Protocol
- 3-Step Decision Logic & Flowchart Guide
- Expanded Desk Cheat Sheet: Statutory Caps, Regulatory Jurisdictions & Supreme Court Precedents
- Submission Mechanics & Panel Allocation Roster (Google Forms)
- Consensus, Majority Voting & Adjudication Protocol
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout_box(doc, title, text, border_color="1F497D", bg_color="F2F5F9"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10.5)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    run_body = p.add_run(text)
    run_body.font.name = "Arial"
    run_body.font.size = Pt(9.5)
    run_body.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    # Add small spacing after callout
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def format_table_headers(table, col_widths, headers, bg_hex="1F497D"):
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_background(hdr_cells[i], bg_hex)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

def format_data_row(row, col_widths, data, is_even=False):
    bg_hex = "F9FBFD" if is_even else "FFFFFF"
    for i, val in enumerate(data):
        cell = row.cells[i]
        cell.text = val
        cell.width = Inches(col_widths[i])
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.0)
            r.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)

def build_guide():
    doc = Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = False
        
        # Header / Footer
        hdr = section.header
        p_hdr = hdr.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Sangguniang Panlungsod Legal Annotation Guide | Ateneo de Davao University")
        r_hdr.font.name = "Arial"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        
        ftr = section.footer
        p_ftr = ftr.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ftr = p_ftr.add_run("CS Undergraduate Thesis — Department of Computer Science — Academic Year 2024–2025")
        r_ftr.font.name = "Arial"
        r_ftr.font.size = Pt(8.5)
        r_ftr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # -------------------------------------------------------------
    # Document Header & Metadata Block
    # -------------------------------------------------------------
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(2)
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("ATENEO DE DAVAO UNIVERSITY\n")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    r_dept = p_inst.add_run("Department of Computer Science — School of Arts and Sciences\nIn Institutional Collaboration with the Sangguniang Panlungsod of Davao City")
    r_dept.font.name = "Arial"
    r_dept.font.size = Pt(9.5)
    r_dept.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(8)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("―" * 45)
    r_div.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    # Main Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("SANGGUNIANG PANLUNGSOD LEGAL ANNOTATION GUIDE\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    r_subtitle = p_title.add_run("Domain Expert Evaluation Guide for Statutory Consistency Auditing of Municipal Ordinances")
    r_subtitle.font.name = "Arial"
    r_subtitle.font.size = Pt(11.5)
    r_subtitle.font.bold = True
    r_subtitle.font.color.rgb = RGBColor(0x2B, 0x5B, 0x84)

    # Research Metadata Table Box
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    col_widths_meta = [2.0, 4.5]
    
    meta_data = [
        ("Research Project:", "“A Coarse-to-Fine Semantic Conflict Detection System for Ex-Ante Davao City Ordinances Using Information Retrieval and Natural Language Inference”"),
        ("Thesis Proponents:", "Ralph Paolo Dulce & Yahyah Odin (Department of Computer Science, AdDU)"),
        ("Advisory & Faculty:", "Mr. Adrian “Ogs” Ablazo (Thesis Adviser) | Ma’am Grace Tacadao (Thesis Professor)"),
        ("Annotator Cohort:", "15 Legal Researchers & Legislative Officers, Sangguniang Panlungsod of Davao City")
    ]
    
    for row_idx, (lbl, val) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = Inches(col_widths_meta[0]), Inches(col_widths_meta[1])
        set_cell_background(cell_lbl, "EDF2F8")
        set_cell_background(cell_val, "F9FBFD")
        set_cell_margins(cell_lbl, top=60, bottom=60, left=100, right=100)
        set_cell_margins(cell_val, top=60, bottom=60, left=100, right=100)
        
        p_l = cell_lbl.paragraphs[0]
        r_l = p_l.add_run(lbl)
        r_l.font.name = "Arial"
        r_l.font.size = Pt(9)
        r_l.font.bold = True
        r_l.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        
        p_v = cell_val.paragraphs[0]
        r_v = p_v.add_run(val)
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9)
        r_v.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    p_spc = doc.add_paragraph()
    p_spc.paragraph_format.space_before = Pt(6)
    p_spc.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 1. Project Purpose & System Context
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. Project Purpose & System Context")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(6)
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.add_run(
        "Welcome! As a legal researcher, legislative committee secretary, or policy staff member of the "
        "Sangguniang Panlungsod (City Council) of Davao City, your professional diligence ensures that every local ordinance "
        "enacted is legally robust, effective, and strictly compliant with higher laws.\n\n"
        "In Philippine local governance, Local Government Units (LGUs) exercise delegated police and regulatory powers "
        "under Republic Act No. 7160 (The Local Government Code of 1991). However, under the well-settled Magtajas Doctrine "
        "(Magtajas v. Pryce Properties, 234 SCRA 255) and related constitutional jurisprudence, a municipal ordinance is strictly "
        "subordinate to national statutes passed by Congress. A local ordinance cannot prohibit what national law permits, "
        "permit what national law prohibits, or alter statutory ceilings.\n\n"
        "Manually auditing proposed ordinances against the entire body of Philippine national statutes (Republic Acts, "
        "Presidential Decrees, and national codes) during the First Reading committee review is an intensive, time-consuming "
        "task prone to cognitive fatigue. To assist legislative staff in this critical oversight function, our undergraduate thesis "
        "develops an automated, artificial intelligence-assisted conflict screening system."
    )

    add_callout_box(
        doc,
        "How the AI System Operates (The Coarse-to-Fine Pipeline)",
        "The system operates in two consecutive stages:\n"
        "1. Stage 1 (Information Retrieval / Coarse Search): Rapidly scans thousands of national statute provisions "
        "using lexical (BM25) and semantic vector search to identify candidate laws governing the subject matter of the draft ordinance.\n"
        "2. Stage 2 (Natural Language Inference / Fine Audit): Employs a specialized Transformer Cross-Encoder model to "
        "conduct a rigorous sentence-level semantic comparison, predicting whether the draft ordinance entails (complies with), "
        "contradicts (violates/exceeds), or remains neutral toward the national statute.\n\n"
        "Your Role: Machine learning models require authoritative domain ground truth to benchmark their accuracy. "
        "Your independent legal judgments on these 350 curated text pairs serve as the gold-standard benchmark against which "
        "the AI system's legal reasoning is evaluated."
    )

    # -------------------------------------------------------------
    # 2. Data Privacy, Anonymity & Research Ethics
    # -------------------------------------------------------------
    h2 = doc.add_heading(level=1)
    r = h2.add_run("2. Data Privacy, Anonymity & Research Ethics (RA 10173)")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_priv = doc.add_paragraph()
    p_priv.paragraph_format.space_after = Pt(6)
    p_priv.paragraph_format.line_spacing = 1.15
    p_priv.add_run(
        "This evaluation strictly adheres to the ethical guidelines of the Ateneo de Davao University Research Ethics Committee "
        "and Republic Act No. 10173 (The Data Privacy Act of 2012). The thesis proponents guarantee full protection of your privacy "
        "and intellectual contributions through the following safeguards:"
    )

    priv_points = [
        ("Complete Evaluator Anonymization: ", "You will be assigned a standardized alphanumeric code (e.g., SP-ANN-01 to SP-ANN-15). Your name, position, or personal identifiers will NEVER appear in the thesis manuscript, public presentations, or published datasets."),
        ("Purely Voluntary Participation: ", "Your participation in this study is completely voluntary. You may pause, request clarifications, or withdraw your participation at any time without any professional or administrative penalty."),
        ("Restricted Academic Purpose: ", "All annotations collected will be used solely for academic validation, statistical agreement modeling (Fleiss' Kappa), and machine learning error analysis within this thesis."),
        ("Aggregated Statistical Reporting: ", "All findings reported in the thesis will be presented strictly in aggregate form (e.g., cross-domain accuracy, panel agreement rates, confusion matrices). No individual response will be singled out or attributed.")
    ]

    for bld, txt in priv_points:
        p_pt = doc.add_paragraph(style='List Bullet')
        p_pt.paragraph_format.space_after = Pt(3)
        r_b = p_pt.add_run(bld)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        p_pt.add_run(txt)

    # -------------------------------------------------------------
    # 3. The Core Task & Single-Premise Pairwise Scope
    # -------------------------------------------------------------
    h3 = doc.add_heading(level=1)
    r = h3.add_run("3. The Core Task & The Single-Premise Rule")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_core = doc.add_paragraph()
    p_core.paragraph_format.space_after = Pt(6)
    p_core.paragraph_format.line_spacing = 1.15
    p_core.add_run(
        "In each assigned item, you will be presented with an isolated pair of legal texts:\n"
        "• National Statute (The Premise): A specific statutory provision extracted from a Philippine Republic Act.\n"
        "• Proposed Local Ordinance (The Hypothesis): A drafted clause from a proposed Davao City municipal ordinance.\n\n"
        "Your guiding evaluative question for every pair is:"
    )

    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(4)
    p_q.paragraph_format.space_after = Pt(6)
    p_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_q = p_q.add_run(
        "“Based strictly and solely on this specific national statute provision,\n"
        "does the draft local ordinance contradict it, entail (comply with) it, or govern an independent matter (neutral)?”"
    )
    r_q.font.name = "Arial"
    r_q.font.size = Pt(10.5)
    r_q.font.bold = True
    r_q.font.italic = True
    r_q.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    add_callout_box(
        doc,
        "CRITICAL EVALUATION RULE: The Strict Single-Premise Constraint",
        "As trained legal professionals, your natural instinct is to evaluate an ordinance against the totality of Philippine "
        "jurisprudence—synthesizing constitutional provisions, the Revised Penal Code, and unlisted administrative circulars.\n\n"
        "⚠️ FOR THIS AI BENCHMARK, YOU MUST SUPPRESS EXTERNAL SYNTHESIS: Judge each pair STRICTLY against the specific "
        "National Statute text provided in that specific row. Even if you know from memory that another unmentioned statute "
        "might prohibit the act, if the provided statutory premise does not address it, the pair must be evaluated solely "
        "on the relationship between the two texts on screen."
    )

    # -------------------------------------------------------------
    # 4. The Three Legal Consistency Choices
    # -------------------------------------------------------------
    h4 = doc.add_heading(level=1)
    r = h4.add_run("4. The Three Evaluation Choices Explained")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    # Contradiction
    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.space_before = Pt(4)
    p_c1.paragraph_format.space_after = Pt(2)
    r_c1 = p_c1.add_run("CHOICE 1: CONTRADICTION (Legal Inconsistency / Ultra Vires Overreach)")
    r_c1.font.bold = True
    r_c1.font.size = Pt(11)
    r_c1.font.color.rgb = RGBColor(0x9C, 0x00, 0x06)

    p_c1_desc = doc.add_paragraph()
    p_c1_desc.paragraph_format.space_after = Pt(4)
    p_c1_desc.paragraph_format.line_spacing = 1.15
    p_c1_desc.add_run(
        "Meaning: The draft local ordinance violates, overrides, diminishes, or exceeds the mandate, bounds, or prohibitions "
        "established in the national statute (Magtajas Doctrine). Common manifestations include:\n"
        "• Exceeding Penalty Ceilings: National law (RA 7160 Sec. 458) limits city penalties to a maximum ₱5,000 fine and 1-year imprisonment. A draft imposing ₱25,000 or permanent property forfeiture is an ultra vires contradiction.\n"
        "• Breaching Temporal Limitations: National law (RA 7581 Sec. 6) limits automatic calamity price freezes to 60 days. A local draft extending the freeze to 180 days is a contradiction.\n"
        "• Diminishing Guaranteed Statutory Rights: National law (RA 11314) guarantees a 20% student fare discount year-round. A local draft reducing it to 10% or suspending it on weekends is a contradiction.\n"
        "• Usurping Exclusive National Regulatory Jurisdiction: National law vests exclusive rate-setting authority in the Energy Regulatory Commission (ERC) and telecommunications franchising in Congress/NTC. A city ordinance attempting to fix electric utility rates or issue local cellular franchises is a contradiction.\n"
        "• Administrative Euphemisms Concealing Statutory Violations: National law (RA 9344) forbids detaining minors for curfew violations. An ordinance requiring apprehended minors to undergo mandatory locked overnight stays (euphemistically termed “protective shelter reflection”) is a contradiction."
    )

    # Entailment
    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.space_before = Pt(4)
    p_c2.paragraph_format.space_after = Pt(2)
    r_c2 = p_c2.add_run("CHOICE 2: ENTAILMENT (Compliant Adoption / Lawful Execution)")
    r_c2.font.bold = True
    r_c2.font.size = Pt(11)
    r_c2.font.color.rgb = RGBColor(0x00, 0x61, 0x00)

    p_c2_desc = doc.add_paragraph()
    p_c2_desc.paragraph_format.space_after = Pt(4)
    p_c2_desc.paragraph_format.line_spacing = 1.15
    p_c2_desc.add_run(
        "Meaning: The draft local ordinance directly executes, adopts, enforces, or lawfully implements the national statute "
        "within the delegated authority of the Local Government Code without deviation or conflict. Common manifestations include:\n"
        "• Direct Statutory Adoption: The local draft mirrors national highway speed limits (80 km/h under RA 4136) or statutory health standards.\n"
        "• Exact Formula Compliance: The draft appropriates at least 20% of its National Tax Allotment (NTA) for local development (RA 7160 Sec. 287) or exactly 5% for disaster funds (RA 10121).\n"
        "• Preserving Mandatory Procedural Clearances: The draft requires sand-and-gravel quarry operators to secure a national DENR Environmental Compliance Certificate (ECC) prior to Mayor's Permit issuance.\n"
        "• Adhering to Diversion Standards: The draft channels apprehended youth to accredited community counseling and family reintegration, faithfully adhering to RA 9344 non-custodial mandates."
    )

    # Neutral
    p_c3 = doc.add_paragraph()
    p_c3.paragraph_format.space_before = Pt(4)
    p_c3.paragraph_format.space_after = Pt(2)
    r_c3 = p_c3.add_run("CHOICE 3: NEUTRAL (Independent / Non-Conflicting Subject Matter)")
    r_c3.font.bold = True
    r_c3.font.size = Pt(11)
    r_c3.font.color.rgb = RGBColor(0x9C, 0x65, 0x00)

    p_c3_desc = doc.add_paragraph()
    p_c3_desc.paragraph_format.space_after = Pt(4)
    p_c3_desc.paragraph_format.line_spacing = 1.15
    p_c3_desc.add_run(
        "Meaning: Both laws may belong to the same broad subject domain (e.g., environmental protection, transportation, or public health), "
        "but they govern distinct administrative obligations with no mutual legal friction, conflict, or direct derivation. Common manifestations include:\n"
        "• Public Health Domain: National law requires immediate hospital reporting of notifiable infectious diseases to DOH, while the local draft sets hygiene protocols for sanitizing public market meat stalls.\n"
        "• Local Governance Domain: National statute establishes juvenile age exemptions (15 years old under RA 9344), while the local draft regulates minimum floor space dimensions for daycare centers (30 sq.m.).\n"
        "• Transportation Domain: National statute establishes licensing criteria for heavy bus operators, while the local draft specifies reflective paint markings for city center bicycle lanes."
    )

    # -------------------------------------------------------------
    # 5. Critical Legal Advisory: Local Autonomy vs. Ultra Vires
    # -------------------------------------------------------------
    h5 = doc.add_heading(level=1)
    r = h5.add_run("5. Legal Advisory: Local Autonomy vs. Ultra Vires (Magtajas Doctrine)")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_adv = doc.add_paragraph()
    p_adv.paragraph_format.space_after = Pt(6)
    p_adv.paragraph_format.line_spacing = 1.15
    p_adv.add_run(
        "As legal advisers to the Sangguniang Panlungsod, evaluators frequently and rightfully champion local autonomy under "
        "Section 16 of the Local Government Code (The General Welfare Clause). To ensure inter-rater consistency across our benchmark, "
        "please apply the established Supreme Court tests governing the constitutional limits of municipal ordinances:"
    )

    add_callout_box(
        doc,
        "The Magtajas Test of Ordinance Validity (Magtajas v. Pryce Properties, 234 SCRA 255)",
        "Under Philippine constitutional jurisprudence, for an ordinance to be valid, it must satisfy six cardinal tests:\n"
        "1. It must not contravene the Constitution or ANY statute passed by Congress.\n"
        "2. It must not be unfair or oppressive.\n"
        "3. It must not be partial or discriminatory.\n"
        "4. It must not prohibit, but may regulate, trade.\n"
        "5. It must be general and consistent with public policy.\n"
        "6. It must not be unreasonable.\n\n"
        "Operational Rule: Local autonomy cannot be invoked to override an express act of Congress. If an ordinance provision "
        "departs from, expands penalties beyond, or removes exceptions guaranteed in a national statute, it is CONTRADICTION (Ultra Vires), "
        "regardless of whether the local policy goal is deemed beneficial for Davao City."
    )

    # -------------------------------------------------------------
    # 6. Difficulty Stratification & Evaluator Calibration
    # -------------------------------------------------------------
    h6 = doc.add_heading(level=1)
    r = h6.add_run("6. Difficulty Stratification & Calibrating Your Attention")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_diff = doc.add_paragraph()
    p_diff.paragraph_format.space_after = Pt(6)
    p_diff.paragraph_format.line_spacing = 1.15
    p_diff.add_run(
        "The 350-pair ground truth corpus is stratified into three distinct complexity tiers. Understanding these tiers will help "
        "you calibrate your mental focus as you review your assigned panel:"
    )

    diff_table = doc.add_table(rows=4, cols=3)
    diff_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    diff_table.autofit = False
    col_w_diff = [1.5, 2.3, 2.7]
    
    diff_headers = ["Tier Level", "Primary Inconsistency Mechanism", "What to Look For"]
    format_table_headers(diff_table, col_w_diff, diff_headers, "1F497D")
    
    diff_rows = [
        ("Tier 1: Surface & Quantitative (30%)", "Direct statutory & numerical cap violations; modal polarity inversions.", "Examine explicit numbers: Fines > ₱5,000, jail terms > 1 year, processing deadlines (3/7/20 days), percentage allocations (<20% dev fund), or 'shall' flipped to 'may'."),
        ("Tier 2: Procedural & Jurisdictional (40%)", "Bypassing mandatory statutory procedures, permits, or agency jurisdiction.", "Examine regulatory authority: Bypassing DENR Environmental Clearance Certificates, usurping LTFRB/LTO transit powers, or erasing senior citizen/student exemptions."),
        ("Tier 3: Latent & Preemption (30%)", "Substantive field preemption, administrative euphemisms, and structural conflict.", "Examine legislative effect: Curfew detention rebranded as 'protective reflection', disguised local taxes on national goods, or local bans on activities licensed by national charters.")
    ]
    
    for idx, r_data in enumerate(diff_rows):
        format_data_row(diff_table.rows[idx+1], col_w_diff, r_data, is_even=(idx%2==1))

    p_spc2 = doc.add_paragraph()
    p_spc2.paragraph_format.space_before = Pt(4)

    # -------------------------------------------------------------
    # 7. Evaluator Confidence Ratings & Optional Notes Protocol
    # -------------------------------------------------------------
    h7 = doc.add_heading(level=1)
    r = h7.add_run("7. Confidence Ratings & Legal Notes Protocol")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_conf = doc.add_paragraph()
    p_conf.paragraph_format.space_after = Pt(4)
    p_conf.paragraph_format.line_spacing = 1.15
    p_conf.add_run(
        "For each text pair in your evaluation form, you will provide three inputs:\n"
        "1. Classification Decision: Contradiction, Entailment, or Neutral.\n"
        "2. Confidence Rating (1 to 3 Stars/Scale): Reflecting how clear-cut the statutory boundary is:\n"
        "   • 3 (High Confidence / Clear-Cut): Explicit statutory ceiling, verbatim compliance, or obvious subject-matter divergence.\n"
        "   • 2 (Moderate Confidence): Clear legal foundation, but requires interpreting administrative phrasing or regulatory scope.\n"
        "   • 1 (Low Confidence / Edge Case): Borderline interpretation, ambiguous statutory wording, or competing statutory interpretations.\n"
        "3. Optional Brief Legal Basis / Note: If you flag a Contradiction or encounter an edge case, please write a concise 1-sentence note "
        "(e.g., “Exceeds ₱5k penalty cap under RA 7160 Sec. 458” or “Encroaches on exclusive LTO driver's license confiscation power”). "
        "These expert notes will provide invaluable qualitative evidence for our thesis discussion and defense."
    )

    # -------------------------------------------------------------
    # 8. Quick 3-Step Decision Flowchart Guide
    # -------------------------------------------------------------
    h8 = doc.add_heading(level=1)
    r = h8.add_run("8. Quick 3-Step Decision Flowchart Guide")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    flow_box = (
        "STEP 1: Check for Legal Friction / Ultra Vires Overreach\n"
        "• Does the local draft prohibit what the national statute expressly permits?\n"
        "• Does it permit what the national statute expressly prohibits?\n"
        "• Does it exceed explicit statutory limits (₱5,000 fine, 1-year jail, 60-day price freeze)?\n"
        "• Does it usurp the exclusive jurisdiction of a national regulatory agency (LTO, ERC, NTC, DOH)?\n"
        "  ➔ IF YES TO ANY: Select CONTRADICTION\n\n"
        "STEP 2: Check for Lawful Statutory Execution / Compliance\n"
        "• Does the local draft directly enforce, implement, or mirror the statutory command?\n"
        "• Does it adopt mandatory statutory formulas (e.g., 20% development fund, 5% disaster fund)?\n"
        "• Does it enforce statutory conditions (e.g., requiring DENR ECC prior to quarry permits)?\n"
        "  ➔ IF YES TO ANY: Select ENTAILMENT\n\n"
        "STEP 3: Check for Independent / Unrelated Subject Matter\n"
        "• Do the two texts govern distinct, non-overlapping administrative duties, even if in the same field?\n"
        "• Does the local draft neither violate, execute, nor depend upon the national statute provided?\n"
        "  ➔ IF YES: Select NEUTRAL"
    )
    add_callout_box(doc, "Systematic 3-Step Evaluation Logic", flow_box, "1F497D", "EDF4F9")

    # -------------------------------------------------------------
    # 9. Desk Cheat Sheet: Statutory Caps & Precedents
    # -------------------------------------------------------------
    h9 = doc.add_heading(level=1)
    r = h9.add_run("9. Desk Cheat Sheet: Statutory Baseline Ceilings & Supreme Court Precedents")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_cs_intro = doc.add_paragraph()
    p_cs_intro.paragraph_format.space_after = Pt(4)
    p_cs_intro.add_run("Keep this quick-reference table open on your desk during evaluation:")

    cs_table = doc.add_table(rows=14, cols=3)
    cs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cs_table.autofit = False
    col_w_cs = [1.6, 2.3, 2.6]
    
    cs_headers = ["Legal Area / Subject", "National Statutory Baseline", "Red Flag in Draft (Contradiction)"]
    format_table_headers(cs_table, col_w_cs, cs_headers, "1F497D")
    
    cheat_sheet_rows = [
        ("Ordinance Penalties", "Max ₱5,000 fine / 1-year imprisonment for cities (RA 7160 Sec. 458(a)(1)(iii)).", "Drafts imposing ₱10,000–₱50,000 fines, 2+ years jail, or real property title forfeiture."),
        ("Calamity Price Control", "Max 60 calendar days automatic price freeze (RA 7581 Sec. 6 / Price Act).", "Drafts establishing 90-day, 120-day, or 180-day local price control periods."),
        ("Processing Time (EODB)", "Max 3 days (Simple), 7 days (Complex), 20 days (Highly Tech) (RA 11032).", "Drafts setting 15-day simple permit reviews or open-ended administrative holds."),
        ("Student Fare Discount", "20% discount on all public land/water transit, year-round (RA 11314).", "Drafts reducing discount to 10%, or suspending discounts during summer breaks/holidays."),
        ("Juvenile Justice", "Minors 15 and below exempt from criminal liability; non-custodial diversion (RA 9344).", "Drafts subjecting minors to lock-up, holding rooms, or curfew jail sentences."),
        ("Real Property Tax", "Basic real property tax capped at max 2% of assessed value for cities (RA 7160 Sec. 233).", "Drafts levying 3% to 5% basic RPT, or adjusting assessment levels without statutory authority."),
        ("Amusement Tax", "Capped at max 30% of gross receipts; school events exempt (RA 7160 Sec. 140).", "Drafts levying 35% amusement tax, or taxing school plays, concerts, and athletic meets."),
        ("Local Development Fund", "At least 20% of annual National Tax Allotment (NTA) (RA 7160 Sec. 287).", "Drafts appropriating less than 20% (e.g., 10%), or diverting development funds to travel/allowances."),
        ("Disaster Fund (LDRRMF)", "At least 5% of estimated regular revenue; 30% allocated to QRF (RA 10121).", "Drafts allocating less than 5% (e.g., 2%), or reverting unexpended QRF to the General Fund early."),
        ("Driver's License Seizure", "Only LTO and officially deputized agents may confiscate driver's licenses (RA 4136).", "Drafts authorizing local traffic enforcers to seize driver's licenses without LTO deputization."),
        ("Air Pollution & Waste", "Open burning of municipal/backyard waste (siga) strictly prohibited (RA 8749 Sec. 20; RA 9003).", "Drafts permitting barangay open burning or unpermitted municipal waste incineration."),
        ("Public Utilities / Telecoms", "Exclusive franchising jurisdiction vested in Congress and NTC (Batangas CATV doctrine).", "Drafts attempting to issue municipal telecom certificates or regulate cell tower radio frequencies."),
        ("Double Regulatory Taxation", "LGUs prohibited from taxing goods/services already taxed by national government (RA 7160 Sec. 133).", "Drafts imposing local inspection fees duplicating national FDA or DTI regulatory charges.")
    ]
    
    for idx, r_data in enumerate(cheat_sheet_rows):
        format_data_row(cs_table.rows[idx+1], col_w_cs, r_data, is_even=(idx%2==1))

    p_spc3 = doc.add_paragraph()
    p_spc3.paragraph_format.space_before = Pt(4)

    # -------------------------------------------------------------
    # 10. Submission Mechanics & Panel Allocation Roster
    # -------------------------------------------------------------
    h10 = doc.add_heading(level=1)
    r = h10.add_run("10. Submission Mechanics & Google Forms Allocation")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(4)
    p_sub.paragraph_format.line_spacing = 1.15
    p_sub.add_run(
        "To ensure balanced evaluator workload and prevent cognitive fatigue, the 350-pair ground truth benchmark "
        "is partitioned into five distinct evaluation sets of exactly 70 pairs each. Evaluators are organized into five panels "
        "with three independent raters per panel (k=3), ensuring statistically rigorous multi-rater consensus:\n"
    )

    roster_table = doc.add_table(rows=6, cols=4)
    roster_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    roster_table.autofit = False
    col_w_ros = [1.2, 1.6, 2.0, 1.7]
    
    ros_headers = ["Panel", "Assigned Set", "Evaluator ID Allocation", "Evaluation Pair IDs"]
    format_table_headers(roster_table, col_w_ros, ros_headers, "1F497D")
    
    roster_data = [
        ("Panel A", "Set A (70 Pairs)", "SP-ANN-01, SP-ANN-02, SP-ANN-03", "GT-001 to GT-070"),
        ("Panel B", "Set B (70 Pairs)", "SP-ANN-04, SP-ANN-05, SP-ANN-06", "GT-071 to GT-140"),
        ("Panel C", "Set C (70 Pairs)", "SP-ANN-07, SP-ANN-08, SP-ANN-09", "GT-141 to GT-210"),
        ("Panel D", "Set D (70 Pairs)", "SP-ANN-10, SP-ANN-11, SP-ANN-12", "GT-211 to GT-280"),
        ("Panel E", "Set E (70 Pairs)", "SP-ANN-13, SP-ANN-14, SP-ANN-15", "GT-281 to GT-350")
    ]
    
    for idx, r_data in enumerate(roster_data):
        format_data_row(roster_table.rows[idx+1], col_w_ros, r_data, is_even=(idx%2==1))

    p_steps = doc.add_paragraph()
    p_steps.paragraph_format.space_before = Pt(6)
    p_steps.paragraph_format.space_after = Pt(4)
    p_steps.paragraph_format.line_spacing = 1.15
    p_steps.add_run("Step-by-Step Google Forms Workflow:\n")

    steps_list = [
        ("Open the Official Link: ", "Access the evaluation Google Form link provided in your panel email or engagement packet."),
        ("Section 1 — Evaluator Verification: ", "Select your assigned Evaluator ID (e.g., SP-ANN-01) and designated Panel/Set from the dropdown. The form will automatically branch into your specific 70-pair block."),
        ("Section 2 — Pair Evaluation: ", "For each item (1 to 70), read the National Statute Premise and Municipal Ordinance Hypothesis, select your classification (Contradiction, Entailment, or Neutral), select your confidence score (1 to 3), and optionally add a brief legal basis note."),
        ("Estimated Time & Pacing: ", "Evaluating 70 items takes approximately 60 to 90 minutes (~1 minute per pair). If logged into your Google account, Google Forms automatically saves draft responses, allowing you to complete the evaluation in 1 or 2 sittings."),
        ("Final Submission: ", "Click SUBMIT at the end of the form. A confirmation screen will display: “Thank you! Your legal evaluation has been recorded.”")
    ]

    for bld, txt in steps_list:
        p_st = doc.add_paragraph(style='List Bullet')
        p_st.paragraph_format.space_after = Pt(3)
        r_b = p_st.add_run(bld)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        p_st.add_run(txt)

    # -------------------------------------------------------------
    # 11. Multi-Rater Consensus & Adjudication Protocol
    # -------------------------------------------------------------
    h11 = doc.add_heading(level=1)
    r = h11.add_run("11. Multi-Rater Consensus & Adjudication Protocol")
    r.font.name = "Arial"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_cons = doc.add_paragraph()
    p_cons.paragraph_format.space_after = Pt(6)
    p_cons.paragraph_format.line_spacing = 1.15
    p_cons.add_run(
        "To alleviate decision fatigue and assure evaluators that individual borderline decisions will not skew the benchmark, "
        "our research implements a rigorous statistical consensus protocol:\n"
        "• Independent Raters: Each pair is evaluated independently by all three legal researchers in the assigned panel without mutual consultation.\n"
        "• Majority Rule Acceptance: If at least 2 out of 3 evaluators agree on the label (e.g., Contradiction - Contradiction - Neutral), the majority judgment is accepted as the final Ground Truth.\n"
        "• Senior Legal Adjudication: In rare instances of a three-way split (1 Contradiction, 1 Entailment, 1 Neutral), the pair is referred to a Senior Legal Researcher / Legislative Counsel for binding qualitative adjudication.\n"
        "• Statistical Validation: Inter-annotator reliability will be formally reported in the thesis using Fleiss' Kappa (κ) and percent agreement."
    )

    # Closing & Gratitude
    p_close = doc.add_paragraph()
    p_close.paragraph_format.space_before = Pt(8)
    p_close.paragraph_format.space_after = Pt(2)
    p_close.paragraph_format.line_spacing = 1.15
    r_cl = p_close.add_run(
        "Daghang Kaayong Salamat! We are deeply grateful for your time, professional dedication, and intellectual contribution "
        "to developing AI-assisted legislative consistency tools for the City of Davao.\n\n"
        "For any inquiries, clarifications, or technical issues during evaluation, please contact:\n"
        "• Ralph Paolo Dulce | Email: rpdulce@addu.edu.ph\n"
        "• Yahyah Odin | Email: ygodin@addu.edu.ph\n"
        "• Mr. Adrian “Ogs” Ablazo (Thesis Adviser) | Department of Computer Science, Ateneo de Davao University"
    )
    r_cl.font.italic = True
    r_cl.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # Save to both target locations
    root_docx = "Legal Annotation Guide (Enhanced).docx"
    docs_docx = os.path.join("docs", "annotation", "Legal_Annotation_Guide_Comprehensive.docx")
    
    os.makedirs(os.path.join("docs", "annotation"), exist_ok=True)
    doc.save(root_docx)
    doc.save(docs_docx)
    
    print(f"Successfully generated:")
    print(f"1. {root_docx} (Filesize: {os.path.getsize(root_docx)} bytes)")
    print(f"2. {docs_docx} (Filesize: {os.path.getsize(docs_docx)} bytes)")

if __name__ == "__main__":
    build_guide()
