import docx

doc_path = r'c:\Users\HP\Documents\GitHub\thesis-repo\docs\annotation\Legal Annotation Guide(1).docx'
doc = docx.Document(doc_path)

# 1. Update P90
p90 = doc.paragraphs[90]
p90.text = ''
r90 = p90.add_run('To ensure balanced evaluator workload and prevent cognitive fatigue, the 350-pair ground truth benchmark is partitioned into five distinct evaluation sets of exactly 70 pairs each. Evaluators are organized into five panels with three independent raters per panel, ensuring statistically rigorous multi-rater consensus (with one member per panel designated as Senior Annotator based on credentials collected during intake):\n')
r90.font.name = 'Inter'

# 2. Update Section 11
# Find target paragraph ('Daghang Kaayong Salamat! ...')
target_idx = None
for i, p in enumerate(doc.paragraphs):
    if 'Daghang Kaayong Salamat!' in p.text:
        target_idx = i
        break

if target_idx is None:
    raise ValueError("Target paragraph 'Daghang Kaayong Salamat!' not found")

target_p = doc.paragraphs[target_idx]

# Helper to add styled run
def add_run(p, text, bold=False, italic=False, font_name='Inter', size=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = font_name
    if size:
        r.font.size = size
    return r

# Remove paragraphs between P101 and target_idx
# (P102 was Senior Legal Adjudication, P103 was Statistical Validation)
p_to_remove = []
for i in range(102, target_idx):
    p_to_remove.append(doc.paragraphs[i])

for p in p_to_remove:
    p._element.getparent().remove(p._element)

# Now insert our new paragraphs before target_p
# 1. Senior Annotator Role & Purpose
p_role = target_p.insert_paragraph_before()
p_role.style = 'normal'
p_role.paragraph_format.left_indent = 457200
p_role.paragraph_format.first_line_indent = -228600
p_role.paragraph_format.space_after = 0
add_run(p_role, 'Senior Annotator Role & Purpose:', bold=True)
add_run(p_role, ' Within each three-member panel (Panels A through E), one evaluator is designated as the Senior Annotator (Panel Lead / Adjudicator). Their core responsibilities include:')

# 1a. Sub-bullet: Binding Qualitative Adjudication
p_role_a = target_p.insert_paragraph_before()
p_role_a.style = 'normal'
p_role_a.paragraph_format.left_indent = 685800
p_role_a.paragraph_format.first_line_indent = -228600
p_role_a.paragraph_format.space_after = 0
add_run(p_role_a, '• Binding Qualitative Adjudication (Tie-Breaking): ', bold=True)
add_run(p_role_a, 'In rare instances of a three-way split (1 Contradiction, 1 Entailment, 1 Neutral) where no 2-out-of-3 majority consensus exists, the pair is referred to the Senior Annotator of that specific panel to perform an authoritative qualitative review against the statutory codebook and determine the final Ground Truth label (following hierarchical annotation frameworks; Gao et al., 2022; Artstein & Poesio, 2008).')

# 1b. Sub-bullet: Resolving Edge Cases & Ambiguities
p_role_b = target_p.insert_paragraph_before()
p_role_b.style = 'normal'
p_role_b.paragraph_format.left_indent = 685800
p_role_b.paragraph_format.first_line_indent = -228600
p_role_b.paragraph_format.space_after = 0
add_run(p_role_b, '• Resolving Edge Cases & Ambiguities: ', bold=True)
add_run(p_role_b, 'When subtle statutory phrasing or administrative euphemisms create uncertainty, the Senior Annotator reviews the pair against the standardized codebook to reconcile whether the discrepancy arose from differing interpretations or an overlooked provisos clause.')

# 1c. Sub-bullet: Panel Point of Contact & Quality Assurance
p_role_c = target_p.insert_paragraph_before()
p_role_c.style = 'normal'
p_role_c.paragraph_format.left_indent = 685800
p_role_c.paragraph_format.first_line_indent = -228600
p_role_c.paragraph_format.space_after = 0
add_run(p_role_c, '• Panel Point of Contact & Quality Assurance: ', bold=True)
add_run(p_role_c, 'Serving as the primary panel contact for procedural questions regarding the annotation guidelines, while ensuring that all initial round ratings remain completely independent.')

# 2. Senior Annotator Selection & Designation Methodology
p_sel = target_p.insert_paragraph_before()
p_sel.style = 'normal'
p_sel.paragraph_format.left_indent = 457200
p_sel.paragraph_format.first_line_indent = -228600
p_sel.paragraph_format.space_after = 0
add_run(p_sel, 'Senior Annotator Selection & Designation Methodology:', bold=True)
add_run(p_sel, ' To ensure an objective, transparent, and credential-backed designation process, Senior Annotators are chosen through a four-step intake protocol established during the initial administrative engagement with the Sangguniang Panlungsod:')

# 2a. Initial Administrative Outreach
p_sel_a = target_p.insert_paragraph_before()
p_sel_a.style = 'normal'
p_sel_a.paragraph_format.left_indent = 685800
p_sel_a.paragraph_format.first_line_indent = -228600
p_sel_a.paragraph_format.space_after = 0
add_run(p_sel_a, '• Initial Administrative Outreach (15 Volunteers): ', bold=True)
add_run(p_sel_a, 'In the proponents\' initial official communication to the Sangguniang Panlungsod transmitting this Legal Annotation Guidebook, the proponents formally request a cohort of 15 volunteer legal researchers and legislative staff across the council offices.')

# 2b. Intake Credential Profiling
p_sel_b = target_p.insert_paragraph_before()
p_sel_b.style = 'normal'
p_sel_b.paragraph_format.left_indent = 685800
p_sel_b.paragraph_format.first_line_indent = -228600
p_sel_b.paragraph_format.space_after = 0
add_run(p_sel_b, '• Intake Credential Profiling: ', bold=True)
add_run(p_sel_b, 'Alongside the volunteer roster, the proponents systematically collect two objective profiling metrics from each participant: (1) Highest Educational Attainment (e.g., Juris Doctor [J.D.], Bachelor of Laws [LL.B.], Master of Laws [LL.M.], Philippine Bar admission, or relevant postgraduate legal degrees); and (2) Years of Professional Experience in legislative drafting, statutory research, ordinance codification, or legal committee review at the Sangguniang Panlungsod.')

# 2c. Objective Panel Stratification & Designation
p_sel_c = target_p.insert_paragraph_before()
p_sel_c.style = 'normal'
p_sel_c.paragraph_format.left_indent = 685800
p_sel_c.paragraph_format.first_line_indent = -228600
p_sel_c.paragraph_format.space_after = 0
add_run(p_sel_c, '• Objective Panel Stratification & Designation: ', bold=True)
add_run(p_sel_c, 'Upon receiving the roster and credentials from the SP Secretariat, the proponents organize the 15 volunteers into the 5 balanced panels (3 members per panel) and objectively designate the evaluator within each panel possessing the highest educational attainment and longest legislative drafting tenure as that panel\'s Senior Annotator (Snow et al., 2008; Zheng et al., 2021; Chalkidis et al., 2022).')

# 2d. Confirmation & Procedural Readiness
p_sel_d = target_p.insert_paragraph_before()
p_sel_d.style = 'normal'
p_sel_d.paragraph_format.left_indent = 685800
p_sel_d.paragraph_format.first_line_indent = -228600
p_sel_d.paragraph_format.space_after = 0
add_run(p_sel_d, '• Confirmation & Procedural Readiness: ', bold=True)
add_run(p_sel_d, 'In the follow-up confirmation email to the SP and the panels, the proponents confirm panel assignments and explicitly communicate who the designated Senior Annotator is for each panel, establishing clear procedural readiness "just in case" qualitative tie-breaking or adjudication is required.')

# 3. Statistical Validation
p_stat = target_p.insert_paragraph_before()
p_stat.style = 'normal'
p_stat.paragraph_format.left_indent = 457200
p_stat.paragraph_format.first_line_indent = -228600
p_stat.paragraph_format.space_after = 76200
add_run(p_stat, 'Statistical Validation: ', bold=True)
add_run(p_stat, 'Inter-annotator reliability will be formally reported in the thesis using Fleiss\' Kappa (κ) and percent agreement across all initial independent ratings prior to adjudication.')

# Save document
doc.save(doc_path)
print('Document updated and saved successfully!')
